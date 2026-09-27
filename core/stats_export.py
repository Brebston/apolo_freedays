"""
Excel-звіт зі статистикою відсутностей.

Не залежить від Django: адмінка збирає дані у прості словники й передає сюди.
Усі підсумки — формули (COUNTIFS) над аркушами «Days» і «Requests», тож якщо
виправити рядок у даних, звіт і графіки перерахуються самі.
"""

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
NAVY = "1F3A5F"
COLORS = {
    "dayoff": "3B6FD8",
    "l4": "E4704A",
    "approved": "2E9E64",
    "pending": "D99A1E",
    "rejected": "D64545",
    "cancelled": "98A2B3",
    "administration": "3B6FD8",
    "accounting": "8E5CD9",
}

LABELS = {
    "en": {
        "title": "Absence statistics", "period": "Period", "generated": "Generated", "project_filter": "Project",
        "all_projects": "All projects", "kpi": "Key figures", "value": "Value",
        "k_requests": "Requests (all statuses)", "k_dayoff_days": "Day-off days", "k_l4_days": "Sick leave (L4) days",
        "k_total_days": "Total absence days", "k_workers": "Workers with absences", "k_approval": "Approval rate (decided day-off requests)",
        "k_pending": "Pending requests", "k_cancelled": "Cancelled by workers", "k_l4_missing": "L4 without a sick note",
        "k_service": "Requests to administration / accounting",
        "by_project": "By project", "project": "Project", "region": "Region", "dayoff_days": "Day-off days",
        "l4_days": "L4 days", "total_days": "Total days", "requests": "Requests",
        "status_pending": "Pending", "status_approved": "Approved", "status_rejected": "Rejected", "status_cancelled": "Cancelled",
        "statuses": "Requests by status", "status": "Status", "count": "Count",
        "monthly": "Absence days by month", "month": "Month", "weekdays": "Absence days by weekday", "weekday": "Weekday",
        "workers": "Absences by worker", "worker": "Worker", "l4_count": "L4 requests", "top_workers": "Top 10 workers by absence days",
        "service": "Requests to administration and accounting", "department": "Department", "total": "Total",
        "dept_administration": "Administration", "dept_accounting": "Accounting",
        "id": "ID", "created": "Created", "type": "Type", "start": "Start", "end": "End", "days": "Days",
        "sick_note": "Sick note", "yes": "yes", "no": "no", "date": "Date", "month_key": "Month key", "weekday_no": "Weekday no.",
        "text": "Text", "type_dayoff": "Day off", "type_l4": "Sick leave (L4)",
        "note_days": "Counts every absence day of requests that are approved or pending. Rejected and cancelled requests are excluded. Source: sheet “Days”.",
        "note_data": "Source data exported from the system. Every figure on the other sheets is a formula over this sheet.",
        "sheet_summary": "Summary", "sheet_monthly": "Monthly", "sheet_weekdays": "Weekdays", "sheet_workers": "Workers",
        "sheet_service": "Admin & accounting", "sheet_requests": "Requests", "sheet_days": "Days", "sheet_service_list": "Service requests",
        "weekday_names": ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"],
        "month_names": ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"],
    },
    "pl": {
        "title": "Statystyka nieobecności", "period": "Okres", "generated": "Wygenerowano", "project_filter": "Projekt",
        "all_projects": "Wszystkie projekty", "kpi": "Kluczowe wskaźniki", "value": "Wartość",
        "k_requests": "Zgłoszenia (wszystkie statusy)", "k_dayoff_days": "Dni wolne", "k_l4_days": "Dni zwolnienia (L4)",
        "k_total_days": "Dni nieobecności łącznie", "k_workers": "Pracownicy z nieobecnościami", "k_approval": "Odsetek zatwierdzeń (rozpatrzone dni wolne)",
        "k_pending": "Zgłoszenia oczekujące", "k_cancelled": "Anulowane przez pracowników", "k_l4_missing": "L4 bez dokumentu",
        "k_service": "Zgłoszenia do administracji / księgowości",
        "by_project": "Według projektu", "project": "Projekt", "region": "Region", "dayoff_days": "Dni wolne",
        "l4_days": "Dni L4", "total_days": "Dni łącznie", "requests": "Zgłoszenia",
        "status_pending": "Oczekuje", "status_approved": "Zatwierdzono", "status_rejected": "Odrzucono", "status_cancelled": "Anulowano",
        "statuses": "Zgłoszenia według statusu", "status": "Status", "count": "Liczba",
        "monthly": "Dni nieobecności według miesięcy", "month": "Miesiąc", "weekdays": "Dni nieobecności według dnia tygodnia", "weekday": "Dzień tygodnia",
        "workers": "Nieobecności według pracowników", "worker": "Pracownik", "l4_count": "Zgłoszenia L4", "top_workers": "Top 10 pracowników wg dni nieobecności",
        "service": "Zgłoszenia do administracji i księgowości", "department": "Dział", "total": "Razem",
        "dept_administration": "Administracja", "dept_accounting": "Księgowość",
        "id": "ID", "created": "Utworzono", "type": "Rodzaj", "start": "Od", "end": "Do", "days": "Dni",
        "sick_note": "Dokument L4", "yes": "tak", "no": "nie", "date": "Data", "month_key": "Klucz miesiąca", "weekday_no": "Nr dnia tygodnia",
        "text": "Treść", "type_dayoff": "Dzień wolny", "type_l4": "Zwolnienie lekarskie (L4)",
        "note_days": "Liczone są wszystkie dni nieobecności ze zgłoszeń zatwierdzonych lub oczekujących. Zgłoszenia odrzucone i anulowane są pominięte. Źródło: arkusz „Dni”.",
        "note_data": "Dane źródłowe wyeksportowane z systemu. Każda liczba na pozostałych arkuszach to formuła na tych danych.",
        "sheet_summary": "Podsumowanie", "sheet_monthly": "Miesiące", "sheet_weekdays": "Dni tygodnia", "sheet_workers": "Pracownicy",
        "sheet_service": "Administracja i księgowość", "sheet_requests": "Zgłoszenia", "sheet_days": "Dni", "sheet_service_list": "Zgłoszenia do działów",
        "weekday_names": ["Poniedziałek", "Wtorek", "Środa", "Czwartek", "Piątek", "Sobota", "Niedziela"],
        "month_names": ["sty", "lut", "mar", "kwi", "maj", "cze", "lip", "sie", "wrz", "paź", "lis", "gru"],
    },
    "uk": {
        "title": "Статистика відсутностей", "period": "Період", "generated": "Сформовано", "project_filter": "Проєкт",
        "all_projects": "Усі проєкти", "kpi": "Ключові показники", "value": "Значення",
        "k_requests": "Зголошення (усі статуси)", "k_dayoff_days": "Дні вихідних", "k_l4_days": "Дні лікарняного (L4)",
        "k_total_days": "Усього днів відсутності", "k_workers": "Працівники з відсутностями", "k_approval": "Частка підтверджень (опрацьовані вихідні)",
        "k_pending": "Очікують рішення", "k_cancelled": "Скасовано працівниками", "k_l4_missing": "L4 без документа",
        "k_service": "Зголошення до адміністрації / бухгалтерії",
        "by_project": "За проєктами", "project": "Проєкт", "region": "Регіон", "dayoff_days": "Дні вихідних",
        "l4_days": "Дні L4", "total_days": "Усього днів", "requests": "Зголошення",
        "status_pending": "Очікує", "status_approved": "Підтверджено", "status_rejected": "Відхилено", "status_cancelled": "Скасовано",
        "statuses": "Зголошення за статусом", "status": "Статус", "count": "Кількість",
        "monthly": "Дні відсутності за місяцями", "month": "Місяць", "weekdays": "Дні відсутності за днями тижня", "weekday": "День тижня",
        "workers": "Відсутності за працівниками", "worker": "Працівник", "l4_count": "Зголошення L4", "top_workers": "Топ-10 працівників за днями відсутності",
        "service": "Зголошення до адміністрації та бухгалтерії", "department": "Відділ", "total": "Разом",
        "dept_administration": "Адміністрація", "dept_accounting": "Бухгалтерія",
        "id": "ID", "created": "Створено", "type": "Тип", "start": "Від", "end": "До", "days": "Днів",
        "sick_note": "Документ L4", "yes": "так", "no": "ні", "date": "Дата", "month_key": "Ключ місяця", "weekday_no": "№ дня тижня",
        "text": "Текст", "type_dayoff": "Вихідний", "type_l4": "Лікарняний (L4)",
        "note_days": "Враховано кожен день відсутності з підтверджених і тих, що очікують, зголошень. Відхилені та скасовані не враховано. Джерело: аркуш «Дні».",
        "note_data": "Вихідні дані, вивантажені із системи. Кожне число на інших аркушах — формула над цими даними.",
        "sheet_summary": "Підсумок", "sheet_monthly": "Місяці", "sheet_weekdays": "Дні тижня", "sheet_workers": "Працівники",
        "sheet_service": "Адміністрація і бухгалтерія", "sheet_requests": "Зголошення", "sheet_days": "Дні", "sheet_service_list": "Запити до відділів",
        "weekday_names": ["Понеділок", "Вівторок", "Середа", "Четвер", "П'ятниця", "Субота", "Неділя"],
        "month_names": ["січ", "лют", "бер", "кві", "тра", "чер", "лип", "сер", "вер", "жов", "лис", "гру"],
    },
}

STATUSES = ["approved", "pending", "rejected", "cancelled"]

_thin = Side(style="thin", color="E4E7EC")
BORDER = Border(bottom=_thin)
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
HEADER_FILL = PatternFill("solid", start_color=NAVY)
BODY_FONT = Font(name=FONT, size=10)
BOLD_FONT = Font(name=FONT, size=10, bold=True)
NOTE_FONT = Font(name=FONT, size=9, italic=True, color="667085")


def _q(sheet_title: str) -> str:
    """Посилання на аркуш у формулі: назва в лапках (може містити пробіли)."""
    return "'" + sheet_title.replace("'", "''") + "'"


def _header(ws, row: int, values: list[str], start_col: int = 1):
    for offset, value in enumerate(values):
        cell = ws.cell(row=row, column=start_col + offset, value=value)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _section_title(ws, row: int, text: str, col: int = 1):
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = Font(name=FONT, size=12, bold=True, color=NAVY)


def _body(cell, bold=False, fmt=None, center=False):
    cell.font = BOLD_FONT if bold else BODY_FONT
    cell.border = BORDER
    if fmt:
        cell.number_format = fmt
    if center:
        cell.alignment = Alignment(horizontal="center")
    return cell


def _widths(ws, widths: dict[str, float]):
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def _style_series(chart, color_keys: list[str]):
    for series, key in zip(chart.series, color_keys):
        series.graphicalProperties.solidFill = COLORS[key]
        series.graphicalProperties.line.solidFill = COLORS[key]


def _month_label(L, key: str) -> str:
    year, month = key.split("-")
    return f"{L['month_names'][int(month) - 1]} {year}"


def build_workbook(*, absences: list[dict], days: list[dict], service_requests: list[dict],
                   date_from: date, date_to: date, project_label: str | None, lang: str = "en") -> bytes:
    """
    absences:  {id, created, worker, project, region, type, status, start, end, days, sick_note}
    days:      {date, worker, project, type}  — лише дні активних (підтверджених/очікуючих) заявок у періоді
    service_requests: {id, created, worker, department, status, text}
    """
    L = LABELS.get(lang, LABELS["en"])
    wb = Workbook()

    ws_sum = wb.active
    ws_sum.title = L["sheet_summary"]
    ws_mon = wb.create_sheet(L["sheet_monthly"])
    ws_wd = wb.create_sheet(L["sheet_weekdays"])
    ws_wrk = wb.create_sheet(L["sheet_workers"])
    ws_svc = wb.create_sheet(L["sheet_service"])
    ws_req = wb.create_sheet(L["sheet_requests"])
    ws_days = wb.create_sheet(L["sheet_days"])
    ws_svc_list = wb.create_sheet(L["sheet_service_list"])

    # ---------- Аркуші даних ----------
    _header(ws_days, 1, [L["date"], L["month_key"], L["weekday_no"], L["project"], L["worker"], L["type"]])
    for i, d in enumerate(sorted(days, key=lambda x: (x["date"], x["worker"])), start=2):
        _body(ws_days.cell(row=i, column=1, value=d["date"]), fmt="DD.MM.YYYY")
        _body(ws_days.cell(row=i, column=2, value=d["date"].strftime("%Y-%m")))
        _body(ws_days.cell(row=i, column=3, value=d["date"].isoweekday()), center=True)
        _body(ws_days.cell(row=i, column=4, value=d["project"]))
        _body(ws_days.cell(row=i, column=5, value=d["worker"]))
        _body(ws_days.cell(row=i, column=6, value=d["type"]))
    days_last = max(2, len(days) + 1)
    ws_days.freeze_panes = "A2"
    ws_days.auto_filter.ref = f"A1:F{days_last}"
    _widths(ws_days, {"A": 12, "B": 12, "C": 12, "D": 22, "E": 28, "F": 10})
    D = _q(ws_days.title)
    d_month, d_wd, d_proj, d_worker, d_type = (f"{D}!${c}$2:${c}${days_last}" for c in "BCDEF")

    req_headers = [L["id"], L["created"], L["worker"], L["project"], L["region"], L["type"], L["status"],
                   L["start"], L["end"], L["days"], L["sick_note"]]
    _header(ws_req, 1, req_headers)
    for i, a in enumerate(sorted(absences, key=lambda x: x["created"]), start=2):
        values = [a["id"], a["created"], a["worker"], a["project"], a["region"], a["type"], a["status"],
                  a["start"], a["end"], a["days"],
                  (L["yes"] if a["sick_note"] else L["no"]) if a["type"] == "l4" else ""]
        for col, value in enumerate(values, start=1):
            fmt = "DD.MM.YYYY HH:MM" if col == 2 else ("DD.MM.YYYY" if col in (8, 9) else None)
            _body(ws_req.cell(row=i, column=col, value=value), fmt=fmt)
    req_last = max(2, len(absences) + 1)
    ws_req.freeze_panes = "A2"
    ws_req.auto_filter.ref = f"A1:K{req_last}"
    _widths(ws_req, {"A": 7, "B": 17, "C": 26, "D": 20, "E": 14, "F": 10, "G": 12, "H": 12, "I": 12, "J": 7, "K": 11})
    R = _q(ws_req.title)
    r_worker, r_proj, r_type, r_status, r_note = (f"{R}!${c}$2:${c}${req_last}" for c in "CDFGK")

    _header(ws_svc_list, 1, [L["id"], L["created"], L["worker"], L["department"], L["status"], L["text"]])
    for i, s in enumerate(sorted(service_requests, key=lambda x: x["created"]), start=2):
        for col, value in enumerate([s["id"], s["created"], s["worker"], s["department"], s["status"], s["text"]], start=1):
            _body(ws_svc_list.cell(row=i, column=col, value=value), fmt="DD.MM.YYYY HH:MM" if col == 2 else None)
    svc_last = max(2, len(service_requests) + 1)
    ws_svc_list.freeze_panes = "A2"
    ws_svc_list.auto_filter.ref = f"A1:F{svc_last}"
    _widths(ws_svc_list, {"A": 7, "B": 17, "C": 26, "D": 16, "E": 12, "F": 70})
    S = _q(ws_svc_list.title)
    s_dept, s_status = (f"{S}!${c}$2:${c}${svc_last}" for c in "DE")

    for ws in (ws_days, ws_req, ws_svc_list):
        last = ws.max_row + 2
        ws.cell(row=last, column=1, value=L["note_data"]).font = NOTE_FONT

    # ---------- Підсумок ----------
    ws_sum["A1"] = L["title"]
    ws_sum["A1"].font = Font(name=FONT, size=18, bold=True, color=NAVY)
    meta = [
        (L["period"], f"{date_from:%d.%m.%Y} — {date_to:%d.%m.%Y}"),
        (L["project_filter"], project_label or L["all_projects"]),
        (L["generated"], datetime.now().strftime("%d.%m.%Y %H:%M")),
    ]
    for offset, (label, value) in enumerate(meta, start=2):
        ws_sum.cell(row=offset, column=1, value=label).font = Font(name=FONT, size=10, color="667085")
        ws_sum.cell(row=offset, column=2, value=value).font = BODY_FONT

    _section_title(ws_sum, 6, L["kpi"])
    kpis = [
        ("k_requests", f"=COUNTA({r_status})-COUNTBLANK({r_status})", "0"),
        ("k_dayoff_days", f'=COUNTIFS({d_type},"dayoff")', "0"),
        ("k_l4_days", f'=COUNTIFS({d_type},"l4")', "0"),
        ("k_total_days", "=B9+B10", "0"),
        ("k_workers", f'=IFERROR(SUMPRODUCT(({d_worker}<>"")/COUNTIF({d_worker},{d_worker}&"")),0)', "0"),
        ("k_approval", f'=IFERROR(COUNTIFS({r_type},"dayoff",{r_status},"approved")/'
                       f'(COUNTIFS({r_type},"dayoff",{r_status},"approved")+COUNTIFS({r_type},"dayoff",{r_status},"rejected")),0)', "0.0%"),
        ("k_pending", f'=COUNTIFS({r_status},"pending")', "0"),
        ("k_cancelled", f'=COUNTIFS({r_status},"cancelled")', "0"),
        ("k_l4_missing", f'=COUNTIFS({r_type},"l4",{r_note},"{L["no"]}",{r_status},"<>cancelled",{r_status},"<>rejected")', "0"),
        ("k_service", f"=COUNTA({s_status})-COUNTBLANK({s_status})", "0"),
    ]
    _header(ws_sum, 7, [L["kpi"], L["value"]])
    for offset, (key, formula, fmt) in enumerate(kpis, start=8):
        _body(ws_sum.cell(row=offset, column=1, value=L[key]))
        _body(ws_sum.cell(row=offset, column=2, value=formula), bold=True, fmt=fmt, center=True)
    kpi_end = 7 + len(kpis)
    ws_sum.cell(row=kpi_end + 1, column=1, value=L["note_days"]).font = NOTE_FONT

    # Таблиця по проєктах
    projects = sorted({a["project"] for a in absences} | {d["project"] for d in days})
    region_of = {a["project"]: a["region"] for a in absences}
    top = kpi_end + 3
    _section_title(ws_sum, top, L["by_project"])
    headers = [L["project"], L["region"], L["dayoff_days"], L["l4_days"], L["total_days"], L["requests"]] + \
              [L[f"status_{s}"] for s in STATUSES]
    _header(ws_sum, top + 1, headers)
    first_proj_row = top + 2
    for i, project in enumerate(projects, start=first_proj_row):
        _body(ws_sum.cell(row=i, column=1, value=project))
        _body(ws_sum.cell(row=i, column=2, value=region_of.get(project, "")))
        _body(ws_sum.cell(row=i, column=3, value=f'=COUNTIFS({d_proj},$A{i},{d_type},"dayoff")'), center=True)
        _body(ws_sum.cell(row=i, column=4, value=f'=COUNTIFS({d_proj},$A{i},{d_type},"l4")'), center=True)
        _body(ws_sum.cell(row=i, column=5, value=f"=C{i}+D{i}"), bold=True, center=True)
        _body(ws_sum.cell(row=i, column=6, value=f"=COUNTIFS({r_proj},$A{i})"), center=True)
        for col, status in enumerate(STATUSES, start=7):
            _body(ws_sum.cell(row=i, column=col, value=f'=COUNTIFS({r_proj},$A{i},{r_status},"{status}")'), center=True)
    last_proj_row = first_proj_row + max(len(projects), 1) - 1
    total_row = last_proj_row + 1
    _body(ws_sum.cell(row=total_row, column=1, value=L["total"]), bold=True)
    for col in range(3, 11):
        letter = get_column_letter(col)
        _body(ws_sum.cell(row=total_row, column=col, value=f"=SUM({letter}{first_proj_row}:{letter}{last_proj_row})"),
              bold=True, center=True)

    # Статуси (дані для кругової діаграми)
    st_top = total_row + 3
    _section_title(ws_sum, st_top, L["statuses"])
    _header(ws_sum, st_top + 1, [L["status"], L["count"]])
    for i, status in enumerate(STATUSES, start=st_top + 2):
        _body(ws_sum.cell(row=i, column=1, value=L[f"status_{status}"]))
        _body(ws_sum.cell(row=i, column=2, value=f'=COUNTIFS({r_status},"{status}")'), center=True)
    _widths(ws_sum, {"A": 44, "B": 16, "C": 13, "D": 11, "E": 12, "F": 12, "G": 13, "H": 12, "I": 12, "J": 12})

    if projects:
        chart = BarChart()
        chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
        chart.title = L["by_project"]
        chart.y_axis.title = L["days"]
        chart.add_data(Reference(ws_sum, min_col=3, max_col=4, min_row=first_proj_row - 1, max_row=last_proj_row), titles_from_data=True)
        chart.set_categories(Reference(ws_sum, min_col=1, min_row=first_proj_row, max_row=last_proj_row))
        _style_series(chart, ["dayoff", "l4"])
        chart.height, chart.width = 8, 16
        ws_sum.add_chart(chart, "L2")

    pie = PieChart()
    pie.title = L["statuses"]
    pie.add_data(Reference(ws_sum, min_col=2, min_row=st_top + 1, max_row=st_top + 5), titles_from_data=True)
    pie.set_categories(Reference(ws_sum, min_col=1, min_row=st_top + 2, max_row=st_top + 5))
    from openpyxl.chart.series import DataPoint
    for idx, status in enumerate(STATUSES):
        point = DataPoint(idx=idx)
        point.graphicalProperties.solidFill = COLORS[status]
        pie.series[0].dPt.append(point)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showVal = False
    pie.dataLabels.showCatName = False
    pie.dataLabels.showSerName = False
    pie.dataLabels.showLegendKey = False
    pie.height, pie.width = 8, 12
    ws_sum.add_chart(pie, "L19")
    ws_sum.freeze_panes = "A6"

    # ---------- Місяці ----------
    months, cursor = [], date(date_from.year, date_from.month, 1)
    while cursor <= date_to:
        months.append(cursor.strftime("%Y-%m"))
        cursor = date(cursor.year + (cursor.month == 12), cursor.month % 12 + 1, 1)
    ws_mon["A1"] = L["monthly"]
    ws_mon["A1"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
    _header(ws_mon, 3, [L["month_key"], L["month"], L["dayoff_days"], L["l4_days"], L["total_days"]])
    for i, key in enumerate(months, start=4):
        _body(ws_mon.cell(row=i, column=1, value=key))
        _body(ws_mon.cell(row=i, column=2, value=_month_label(L, key)))
        _body(ws_mon.cell(row=i, column=3, value=f'=COUNTIFS({d_month},$A{i},{d_type},"dayoff")'), center=True)
        _body(ws_mon.cell(row=i, column=4, value=f'=COUNTIFS({d_month},$A{i},{d_type},"l4")'), center=True)
        _body(ws_mon.cell(row=i, column=5, value=f"=C{i}+D{i}"), bold=True, center=True)
    mon_last = 3 + len(months)
    ws_mon.column_dimensions["A"].hidden = True
    _widths(ws_mon, {"B": 14, "C": 14, "D": 12, "E": 14})
    ws_mon.cell(row=mon_last + 2, column=2, value=L["note_days"]).font = NOTE_FONT
    chart = BarChart()
    chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
    chart.title = L["monthly"]
    chart.y_axis.title = L["days"]
    chart.add_data(Reference(ws_mon, min_col=3, max_col=4, min_row=3, max_row=mon_last), titles_from_data=True)
    chart.set_categories(Reference(ws_mon, min_col=2, min_row=4, max_row=mon_last))
    _style_series(chart, ["dayoff", "l4"])
    chart.height, chart.width = 9, 20
    ws_mon.add_chart(chart, "G3")

    # ---------- Дні тижня ----------
    ws_wd["A1"] = L["weekdays"]
    ws_wd["A1"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
    _header(ws_wd, 3, [L["weekday_no"], L["weekday"], L["dayoff_days"], L["l4_days"], L["total_days"]])
    for i, name in enumerate(L["weekday_names"], start=4):
        number = i - 3
        _body(ws_wd.cell(row=i, column=1, value=number), center=True)
        _body(ws_wd.cell(row=i, column=2, value=name))
        _body(ws_wd.cell(row=i, column=3, value=f'=COUNTIFS({d_wd},$A{i},{d_type},"dayoff")'), center=True)
        _body(ws_wd.cell(row=i, column=4, value=f'=COUNTIFS({d_wd},$A{i},{d_type},"l4")'), center=True)
        _body(ws_wd.cell(row=i, column=5, value=f"=C{i}+D{i}"), bold=True, center=True)
    ws_wd.column_dimensions["A"].hidden = True
    _widths(ws_wd, {"B": 16, "C": 14, "D": 12, "E": 14})
    ws_wd.cell(row=12, column=2, value=L["note_days"]).font = NOTE_FONT
    chart = BarChart()
    chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
    chart.title = L["weekdays"]
    chart.add_data(Reference(ws_wd, min_col=3, max_col=4, min_row=3, max_row=10), titles_from_data=True)
    chart.set_categories(Reference(ws_wd, min_col=2, min_row=4, max_row=10))
    _style_series(chart, ["dayoff", "l4"])
    chart.height, chart.width = 9, 18
    ws_wd.add_chart(chart, "G3")

    # ---------- Працівники ----------
    per_worker = {}
    for d in days:
        per_worker[d["worker"]] = per_worker.get(d["worker"], 0) + 1
    workers = sorted({a["worker"] for a in absences}, key=lambda w: (-per_worker.get(w, 0), w))
    ws_wrk["A1"] = L["workers"]
    ws_wrk["A1"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
    _header(ws_wrk, 3, [L["worker"], L["dayoff_days"], L["l4_days"], L["total_days"], L["l4_count"], L["requests"]])
    for i, worker in enumerate(workers, start=4):
        _body(ws_wrk.cell(row=i, column=1, value=worker))
        _body(ws_wrk.cell(row=i, column=2, value=f'=COUNTIFS({d_worker},$A{i},{d_type},"dayoff")'), center=True)
        _body(ws_wrk.cell(row=i, column=3, value=f'=COUNTIFS({d_worker},$A{i},{d_type},"l4")'), center=True)
        _body(ws_wrk.cell(row=i, column=4, value=f"=B{i}+C{i}"), bold=True, center=True)
        _body(ws_wrk.cell(row=i, column=5, value=f'=COUNTIFS({r_worker},$A{i},{r_type},"l4",{r_status},"<>cancelled",{r_status},"<>rejected")'), center=True)
        _body(ws_wrk.cell(row=i, column=6, value=f"=COUNTIFS({r_worker},$A{i})"), center=True)
    ws_wrk.freeze_panes = "A4"
    _widths(ws_wrk, {"A": 28, "B": 13, "C": 11, "D": 13, "E": 12, "F": 12})
    if workers:
        top_n = min(10, len(workers))
        chart = BarChart()
        chart.type, chart.grouping, chart.overlap = "bar", "stacked", 100
        chart.title = L["top_workers"]
        chart.add_data(Reference(ws_wrk, min_col=2, max_col=3, min_row=3, max_row=3 + top_n), titles_from_data=True)
        chart.set_categories(Reference(ws_wrk, min_col=1, min_row=4, max_row=3 + top_n))
        chart.y_axis.scaling.orientation = "minMax"
        chart.x_axis.scaling.orientation = "maxMin"  # найбільший зверху
        _style_series(chart, ["dayoff", "l4"])
        chart.height, chart.width = 10, 18
        ws_wrk.add_chart(chart, "H3")

    # ---------- Адміністрація і бухгалтерія ----------
    ws_svc["A1"] = L["service"]
    ws_svc["A1"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
    _header(ws_svc, 3, [L["department"]] + [L[f"status_{s}"] for s in STATUSES] + [L["total"]])
    for i, dept in enumerate(["administration", "accounting"], start=4):
        _body(ws_svc.cell(row=i, column=1, value=L[f"dept_{dept}"]))
        for col, status in enumerate(STATUSES, start=2):
            _body(ws_svc.cell(row=i, column=col, value=f'=COUNTIFS({s_dept},"{dept}",{s_status},"{status}")'), center=True)
        _body(ws_svc.cell(row=i, column=6, value=f"=SUM(B{i}:E{i})"), bold=True, center=True)
    _widths(ws_svc, {"A": 22, "B": 13, "C": 12, "D": 12, "E": 12, "F": 10})
    chart = BarChart()
    chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
    chart.title = L["service"]
    chart.add_data(Reference(ws_svc, min_col=2, max_col=5, min_row=3, max_row=5), titles_from_data=True)
    chart.set_categories(Reference(ws_svc, min_col=1, min_row=4, max_row=5))
    _style_series(chart, STATUSES)
    chart.height, chart.width = 8, 14
    ws_svc.add_chart(chart, "H3")

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
